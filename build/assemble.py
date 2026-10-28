#!/usr/bin/env python
"""Builds deck.pptx (artwork pictures + editable text boxes) from build/art and build/meta."""
import glob, io, json, os, re, sys
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
B = os.path.join(ROOT, 'build')
PX = 6350  # EMU per design pixel (1920 px = 12192000 EMU)


def color(s):
    m = re.match(r'rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)', s or '')
    return RGBColor(int(float(m.group(1))), int(float(m.group(2))), int(float(m.group(3)))) if m else RGBColor(242, 245, 250)


def add_text(slide, it):
    tb = slide.shapes.add_textbox(Emu(int(it['x'] * PX)), Emu(int(it['y'] * PX)), Emu(int((it['w'] + 6) * PX)), Emu(int(it['h'] * PX)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.line_spacing = Pt(it['lh'] * 0.5)
    p.alignment = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}.get(it.get('align'), PP_ALIGN.LEFT)
    for r in it['runs']:
        if r['t'] == '\n':
            p.add_line_break()
            continue
        if not r['t']:
            continue
        run = p.add_run()
        run.text = r['t']
        f = run.font
        f.name = 'Consolas' if r.get('mono') else 'Segoe UI Semibold' if r.get('bold') else 'Segoe UI'
        f.size = Pt(it['size'] * 0.5 * (0.9 if r.get('mono') else 1))
        f.bold = False
        f.color.rgb = color(r.get('color'))
        if it.get('spacing'):
            run._r.get_or_add_rPr().set('spc', str(int(it['spacing'] * 0.5 * 100)))


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    nums = sorted(int(os.path.basename(p)[:2]) for p in glob.glob(os.path.join(B, 'art', '*.png')))
    for n in nums:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(os.path.join(B, 'art', '%02d.png' % n), 0, 0, prs.slide_width, prs.slide_height)
        mp = os.path.join(B, 'meta', '%02d.json' % n)
        items = json.load(io.open(mp, encoding='utf8')) if os.path.exists(mp) else []
        for it in items:
            if it['runs']:
                add_text(s, it)
        vid = os.path.join(ROOT, 'videos', '%02d.mp4' % n)
        page = os.path.join(B, 'html', 'auto', '%02d.html' % n)
        mv = re.search(r'data-video="([^"]+)"', io.open(page, encoding='utf8').read()) if os.path.exists(page) else None
        if mv:
            vid = os.path.join(ROOT, 'videos', mv.group(1))
        if os.path.exists(vid):
            s.shapes.add_movie(vid, Emu(112 * PX), Emu(250 * PX), Emu(1280 * PX), Emu(720 * PX),
                               poster_frame_image=os.path.join(B, 'art', '%02d.png' % n), mime_type='video/mp4')
    prs.core_properties.title = 'The review loop'
    prs.core_properties.author = 'Kenneth Rohde Christiansen'
    out = os.path.join(ROOT, 'deck.pptx')
    prs.save(out)
    print('wrote', out, len(nums), 'slides', round(os.path.getsize(out) / 1e6, 1), 'MB')


main()
